"""Generate a journal-style editable Word document from the review Markdown."""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "工业狭小空间几何测量方法研究进展.md"
TARGET = ROOT / "工业狭小空间几何测量方法研究进展.docx"


def font(run, east_asia="宋体", latin="Times New Roman", size=10.5, bold=None):
    run.font.name = latin
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    if bold is not None:
        run.bold = bold


def clean(text):
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    return text.replace("\\(", "").replace("\\)", "")


def inline(paragraph, text, size=10.5):
    position = 0
    for match in re.finditer(r"\*\*(.*?)\*\*|\*(.*?)\*", text):
        if match.start() > position:
            font(paragraph.add_run(text[position : match.start()]), size=size)
        run = paragraph.add_run(match.group(1) or match.group(2))
        font(run, size=size, bold=match.group(1) is not None)
        run.italic = match.group(2) is not None
        position = match.end()
    if position < len(text):
        font(paragraph.add_run(text[position:]), size=size)


def body_format(paragraph, indent=True):
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    if indent:
        fmt.first_line_indent = Cm(0.74)


def heading(document, text, level):
    paragraph = document.add_paragraph(style=f"Heading {min(level, 3)}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10 if level > 1 else 0)
    paragraph.paragraph_format.space_after = Pt(6)
    size = {1: 18, 2: 14, 3: 12}[level]
    if level == 1:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(paragraph.add_run(clean(text)), east_asia="黑体", size=size, bold=True)


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def table(document, lines):
    values = [[clean(value.strip()) for value in line.strip("|").split("|")] for line in lines]
    if len(values) > 1 and all(re.fullmatch(r":?-{3,}:?", value) for value in values[1]):
        values.pop(1)
    output = document.add_table(rows=len(values), cols=len(values[0]))
    output.style = "Table Grid"
    output.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_index, row_values in enumerate(values):
        for column_index, value in enumerate(row_values):
            cell = output.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            )
            run = paragraph.add_run(value)
            font(run, size=8.5, bold=row_index == 0)
            if row_index == 0:
                shade(cell, "D9EAF7")
    document.add_paragraph()


def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for kind, value in (("begin", None), (None, " PAGE "), ("separate", None), ("end", None)):
        if value is not None:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = value
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        run._r.append(node)
    font(run, size=9)


def main():
    document = Document()
    document.core_properties.title = "工业狭小空间几何测量方法研究进展：系统性综述"
    document.core_properties.subject = "工业狭小空间几何测量系统性综述"
    document.core_properties.keywords = "狭小空间；几何测量；系统性综述；深孔；内窥视觉"

    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin, section.right_margin = Cm(2.8), Cm(2.6)
    page_number(section.footer.paragraphs[0])

    style = document.styles["Normal"]
    style.font.name, style.font.size = "Times New Roman", Pt(10.5)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        text = lines[index].strip()
        if not text or text == "---":
            index += 1
            continue
        if text.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(lines[index].strip())
                index += 1
            table(document, rows)
            continue
        if text == "\\[":
            formula = []
            index += 1
            while index < len(lines) and lines[index].strip() != "\\]":
                formula.append(lines[index].strip())
                index += 1
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(clean(" ".join(formula)))
            font(run, east_asia="Cambria Math", latin="Cambria Math", size=10.5)
            index += 1
            continue
        match = re.match(r"^(#{1,3})\s+(.+)$", text)
        if match:
            heading(document, match.group(2), len(match.group(1)))
            index += 1
            continue
        match = re.match(r"^\[(\d+)\]\s+(.+)$", text)
        if match:
            paragraph = document.add_paragraph()
            inline(paragraph, text, size=9)
            paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.first_line_indent = Cm(-0.74)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_after = Pt(3)
            index += 1
            continue
        match = re.match(r"^(\d+)\.\s+(.+)$", text)
        if match:
            paragraph = document.add_paragraph(style="List Number")
            inline(paragraph, match.group(2))
            body_format(paragraph, indent=False)
            paragraph.paragraph_format.left_indent = Cm(0.74)
            index += 1
            continue
        paragraph = document.add_paragraph()
        inline(paragraph, text)
        body_format(
            paragraph,
            indent=not (text.startswith("**关键词") or text.startswith("**Keywords")),
        )
        index += 1

    document.save(TARGET)
    print(TARGET)


if __name__ == "__main__":
    main()
